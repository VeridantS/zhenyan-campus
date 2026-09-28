from pathlib import Path
import html
import re
import shutil
import zipfile
import math
import hashlib
import argparse

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips, RGBColor
import pypdfium2 as pdfium

HERE = Path(__file__).resolve().parent
PACKED = (HERE / 'content3.md').exists()
ROOT = HERE if PACKED else HERE.parent
OUT = HERE if PACKED else ROOT / "output" / "pdf"
FIG = OUT / "figures"
STEM = "软件工程2班-24111302085-沈振国-实验3-软件架构与界面设计"
REPORT_MD = OUT / f"{STEM}.md"
REPORT_PDF = OUT / f"{STEM}.pdf"
REPORT_DOCX = OUT / f"{STEM}.docx"
QA = HERE / '.preview' if PACKED else ROOT / "tmp" / "report3-final"
pdfmetrics.registerFont(TTFont("CN", "C:/Windows/Fonts/simhei.ttf"))
INK = colors.HexColor("#263b48")
BLUE = colors.HexColor("#eaf0f5")


def diagram(name, title, height):
    c = canvas.Canvas(str(FIG / (name + '.pdf')), pagesize=(680, height))
    c.setFillColor(colors.white)
    c.rect(0, 0, 680, height, fill=1, stroke=0)
    c.setFillColor(INK)
    c.setFont("CN", 15)
    c.drawString(20, height - 28, title)
    return c


def box(c, x, y, w, h, label, font=11):
    c.setFillColor(BLUE); c.setStrokeColor(colors.HexColor('#b8c7d3'))
    c.roundRect(x, y, w, h, 7, fill=1, stroke=1)
    c.setFillColor(INK); c.setFont('CN', font)
    lines = label.split('\n')
    for i, line in enumerate(lines):
        assert pdfmetrics.stringWidth(line, 'CN', font) < w - 12, line
        c.drawCentredString(x+w/2, y+h/2+(len(lines)-1)*7-i*14-4, line)


def arrow(c, points, label='', at=None):
    c.setStrokeColor(colors.HexColor('#758b9b')); c.setLineWidth(1)
    for start, end in zip(points, points[1:]): c.line(*start, *end)
    (sx, sy), (ex, ey) = points[-2:]
    angle = math.atan2(ey-sy, ex-sx)
    for off in (-.45, .45):
        c.line(ex, ey, ex-7*math.cos(angle+off), ey-7*math.sin(angle+off))
    if label:
        tx, ty = at if at else ((sx+ex)/2, (sy+ey)/2+5)
        c.setFont('CN', 9); c.setFillColor(INK); c.drawCentredString(tx, ty, label)


def save_diagram(c, name):
    c.save()
    doc = pdfium.PdfDocument(str(FIG / (name + '.pdf')))
    page = doc[0]; page.render(scale=2).to_pil().save(FIG / (name + '.png')); page.close()
    doc.close()


def draw_architecture():
    c = diagram('architecture', '图3-2 单体架构：主程序组织调用与数据保存', 435)
    box(c, 40, 310, 260, 57, '浏览器前端\n页面 / 导航 / 输入与反馈')
    box(c, 40, 190, 260, 70, 'app.js + Express\n会话 / CSRF / 组织范围 / 业务SQL')
    box(c, 40, 45, 260, 65, 'db.js + SQLite\n初始化 / 持久化 / 迁移')
    box(c, 440, 300, 215, 53, 'checker.js\n规则与正文匹配')
    box(c, 440, 190, 215, 70, 'link-preview.js / ai.js\n公网网页 / 云模型HTTPS调用\n数据返回主程序', 10)
    box(c, 440, 80, 215, 53, 'auth.js\n密码散列 / 令牌 / 字段投影')
    arrow(c, [(170,310),(170,260)], 'JSON请求与响应', (234,281))
    arrow(c, [(170,190),(170,110)], '参数化SQL读写', (237,148))
    arrow(c, [(300,245),(360,245),(360,326),(440,326)], '筛选后的通知', (391,338))
    arrow(c, [(300,225),(440,225)], '受控外部调用', (367,237))
    arrow(c, [(300,205),(335,205),(335,106),(440,106)], '安全工具', (389,118))
    c.setFont('CN',9); c.setFillColor(INK)
    c.drawString(40,20,'箭头表示调用方向；返回值沿调用返回。辅助模块不直接访问数据库。')
    save_diagram(c, 'architecture')


def draw_flow():
    c = diagram('flow', '图3-4 检测、失败返回与人工复核流程', 640)
    box(c,25,533,270,46,'学生输入文字 / 链接 / 可选图片')
    box(c,25,443,270,46,'校验会话、CSRF、输入与限流')
    box(c,25,343,270,58,'可选网页 / AI + 基础检测\n正常数据进入规则与正文匹配')
    box(c,25,243,270,55,'保存本人检测记录\n展示风险依据与不确定性')
    box(c,25,133,270,52,'学生决定是否申请人工复核')
    box(c,25,35,270,48,'暂不申请：结束\n记录仅在本人历史中可见')
    box(c,405,443,250,55,'校验失败 / 429：本次终止\n重新登录 / 修改输入 / 稍后重试',10)
    box(c,405,343,250,58,'外部失败：提示原因并降级\n保留可用文字或链接检测\n图片未识别则要求补充文字',10)
    box(c,405,133,250,52,'POST /reports 主动申请\n保存指定记录，状态open')
    box(c,405,35,250,48,'授权老师核实 / PATCH /reports/:id\n填写reply，reviewing或closed',10)
    box(c,405,243,250,55,'GET /reports 查看进度与答复\n未填写reply时显示待处理',10)
    arrow(c,[(160,533),(160,489)])
    arrow(c,[(160,443),(160,401)],'通过',(183,419))
    arrow(c,[(295,466),(405,466)],'失败',(350,478))
    arrow(c,[(530,498),(530,556),(295,556)],'修正后重新提交',(413,568))
    arrow(c,[(160,343),(160,298)],'正常',(184,316))
    arrow(c,[(295,372),(405,372)],'外部失败',(350,384))
    arrow(c,[(430,343),(430,323),(285,323),(285,298)],'可用内容继续；无内容不作安全判断',(470,309))
    arrow(c,[(160,243),(160,185)])
    arrow(c,[(160,133),(160,83)],'否',(181,106))
    arrow(c,[(295,159),(405,159)],'是 / 用户确认',(351,171))
    arrow(c,[(530,133),(530,83)],'范围外404，数据不变',(590,103))
    arrow(c,[(655,59),(670,59),(670,270),(655,270)])
    save_diagram(c,'flow')


def plain(value):
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", value)
    return value.replace("**", "").replace("`", "")


def parse_blocks(markdown):
    lines = markdown.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.strip() == "<!-- PAGEBREAK -->":
            blocks.append(("pagebreak", ""))
            i += 1
            continue
        if line.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(buf)))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                current = lines[i]
                if not re.match(r"^\|[\s:|\-]+\|$", current):
                    rows.append([plain(cell.strip()) for cell in current.strip("|").split("|")])
                i += 1
            blocks.append(("table", rows))
            continue
        if line.startswith("!["):
            match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", line)
            blocks.append(("image", (match.group(1), match.group(2))))
            i += 1
            continue
        if line.startswith("#"):
            match = re.match(r"(#+)\s+(.*)", line)
            blocks.append((f"h{len(match.group(1))}", plain(match.group(2))))
            i += 1
            continue
        blocks.append(("p", plain(line.lstrip("> ").strip())))
        i += 1
    return blocks


STYLES = {
    "p": ParagraphStyle("body", fontName="CN", fontSize=10.1, leading=15.5, spaceAfter=7, wordWrap="CJK"),
    "h1": ParagraphStyle("title", fontName="CN", fontSize=21, leading=28, spaceAfter=18),
    "h2": ParagraphStyle("section", fontName="CN", fontSize=15, leading=21, spaceBefore=11, spaceAfter=9, keepWithNext=True),
    "h3": ParagraphStyle("subsection", fontName="CN", fontSize=12, leading=17, spaceBefore=9, spaceAfter=7, keepWithNext=True),
    "cell": ParagraphStyle("cell", fontName="CN", fontSize=8.2, leading=11.8, wordWrap="CJK"),
    "code": ParagraphStyle("code", fontName="CN", fontSize=8.2, leading=11.3, wordWrap="CJK"),
}


def ptext(value, style="p"):
    p = Paragraph(html.escape(value).replace("\n", "<br/>").replace("  ", "&nbsp;&nbsp;"), STYLES[style])
    if style == 'p' and re.match(r'^(表3-|代码3-)', value): p.keepWithNext = True
    return p


def build_pdf(blocks):
    story = []
    for kind, data in blocks:
        if kind == "pagebreak":
            story.append(PageBreak())
        elif kind in STYLES:
            story.append(ptext(data, kind))
        elif kind == "code":
            story.append(ptext(data, "code"))
            story.append(Spacer(1, 5))
        elif kind == "image":
            image_path = OUT / data[1]
            from PIL import Image as PILImage
            width, height = PILImage.open(image_path).size
            max_width = 480
            scale = min(1, max_width / width)
            story.extend([Image(str(image_path), width=width * scale, height=height * scale), Spacer(1, 7)])
        elif kind == "table":
            if not data:
                continue
            count = len(data[0])
            total = 480
            if count == 2:
                widths = [125, total - 125]
            elif count == 3:
                widths = [85, 185, total - 270]
            elif count == 4:
                widths = [72, 130, 140, total - 342]
            else:
                widths = [60] + [(total - 60) / (count - 1)] * (count - 1)
            table = Table([[ptext(cell, "cell") for cell in row] for row in data], colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), BLUE),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d9d9d9")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([table, Spacer(1, 9)])

    def footer(c, doc):
        c.setFont("CN", 8)
        c.setFillColor(colors.HexColor("#66727b"))
        c.drawString(55, 28, "真言校园  |  沈振国 24111302085")
        c.drawRightString(540, 28, str(doc.page))

    SimpleDocTemplate(str(REPORT_PDF), pagesize=A4, leftMargin=55, rightMargin=55, topMargin=45, bottomMargin=48,
                      title=STEM, author="沈振国").build(story, onFirstPage=footer, onLaterPages=footer)


def set_cell(cell, value, header=False):
    cell.text = value
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.1
        for run in paragraph.runs:
            run.font.name = "宋体"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
            run.font.size = Pt(8.5)
            run.bold = header
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    tc_pr.append(borders)
    if header:
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "EAF0F5")
        tc_pr.append(shade)


def build_docx(blocks):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)
    doc.core_properties.author = "沈振国"; doc.core_properties.title = STEM
    for name, size in (("Normal", 10.5), ("Title", 21), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "宋体"; style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size); style.paragraph_format.line_spacing = 1.25; style.paragraph_format.space_after = Pt(7)
        style.font.color.rgb = RGBColor(0,0,0)
        if name.startswith("Heading"): style.paragraph_format.keep_with_next = True
    for kind, data in blocks:
        if kind == "pagebreak":
            doc.add_page_break()
        elif kind.startswith("h"):
            doc.add_paragraph(data, style={"h1": "Title", "h2": "Heading 1", "h3": "Heading 2"}.get(kind, "Heading 3"))
        elif kind in ("p", "code"):
            paragraph = doc.add_paragraph(data[2:], style='List Bullet') if data.startswith('- ') else doc.add_paragraph(data)
            if re.match(r'^(表3-|代码3-)', data): paragraph.paragraph_format.keep_with_next = True
            if kind == "code":
                for run in paragraph.runs: run.font.size = Pt(8.5)
        elif kind == "image":
            paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture = paragraph.add_run().add_picture(str(OUT / data[1]), width=Cm(15.5))
            picture._inline.docPr.set("descr", data[0])
        elif kind == "table":
            if not data: continue
            table = doc.add_table(rows=0, cols=len(data[0])); table.autofit = False
            for row_index, row in enumerate(data):
                cells = table.add_row().cells
                tr_pr = table.rows[-1]._tr.get_or_add_trPr(); tr_pr.append(OxmlElement("w:cantSplit"))
                if row_index == 0: tr_pr.append(OxmlElement("w:tblHeader"))
                for cell, value in zip(cells, row): set_cell(cell, value, row_index == 0)
            count = len(data[0]); total = int(Cm(17).twips)
            weights = {2:[125,355],3:[85,185,210],4:[72,130,140,138],5:[60,110,110,75,125]}[count]
            widths = [round(total*w/sum(weights)) for w in weights]; widths[-1] += total-sum(widths)
            props = table._tbl.tblPr
            tw = props.find(qn('w:tblW')); tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'),str(total))
            ind=OxmlElement('w:tblInd'); ind.set(qn('w:type'),'dxa');ind.set(qn('w:w'),'120');props.append(ind)
            margins=OxmlElement('w:tblCellMar')
            for side,val in [('top',70),('bottom',70),('start',120),('end',120)]:
                el=OxmlElement('w:'+side);el.set(qn('w:type'),'dxa');el.set(qn('w:w'),str(val));margins.append(el)
            props.append(margins)
            for col,width in zip(table.columns,widths): col.width=Twips(width)
            for row in table.rows:
                for cell,width in zip(row.cells,widths):cell.width=Twips(width)
            doc.add_paragraph()
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("真言校园 | 沈振国 24111302085 | ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    doc.save(REPORT_DOCX)


def prepare_markdown():
    content = (OUT / "content3.md").read_text(encoding="utf-8")
    old = REPORT_MD.read_text(encoding="utf-8") if REPORT_MD.exists() else ""
    head = old[:old.index("## 3.")] if "## 3." in old else "# 实验3 软件架构与界面设计\n\n"
    head = re.sub(r'报告(?:整理|修订)日期：[^\n]+', '报告修订日期：2026年9月28日　设计版本：V1.2', head)
    report = head + "## 3. 操作记录\n\n" + content
    REPORT_MD.write_text(report, encoding="utf-8")
    return report


def make_upload_package():
    package = OUT / "report3-upload"
    target = package / "reports" / "experiment3"; target.mkdir(parents=True, exist_ok=True)
    for source in (REPORT_PDF, REPORT_DOCX, REPORT_MD, OUT/'content3.md'): shutil.copy2(source, target / source.name)
    figure_target = target / "figures"; figure_target.mkdir(exist_ok=True)
    names = ["context", "architecture", "data", "flow", "navigation", "prototype-login", "prototype-check", "prototype-result"]
    for name in names:
        for suffix in (".png", ".pdf"):
            source = FIG / f"{name}{suffix}"
            if source.exists(): shutil.copy2(source, figure_target / source.name)
    shutil.copy2(Path(__file__), target / "build-report3.py")
    (target/'requirements.txt').write_text('reportlab==4.4.9\npython-docx==1.2.0\npypdfium2==5.12.1\nPillow==12.2.0\n',encoding='utf-8')
    (target / "README.md").write_text('''# 实验3 软件架构与界面设计 V1.2

作者：沈振国，软件工程2班，24111302085。修订日期：2026年9月28日。

## 文件说明

- 同名PDF为已检查排版的阅读版；Word为可编辑版，尚未完成独立页面渲染验证。
- 同名Markdown保留完整要求和正文；content3.md是正文编辑源。
- figures含8组设计图，各有PNG与PDF；Markdown中的相对图片路径可在GitHub展示。
- build-report3.py可在本目录重建报告，并生成修正后的架构图和流程图；其余6张图复用所附图件。
- requirements.txt列生成依赖；manifest.sha256记录文件完整性（不含清单自身）。

## 网页上传

1. 打开 https://github.com/VeridantS/zhenyan-campus/tree/main/reports 。
2. 选择 Add file → Upload files。
3. 从本地 reports 文件夹把 experiment3 文件夹整体拖入网页，保留 figures 子目录；不要只上传zip。
4. 提交说明填“添加实验3架构与界面设计V1.2”，点击 Commit changes。
5. 打开刚才的提交，复制浏览器地址；这才是实验3的真实提交证据。

## 独立重建（可选，无须运行也能阅读和上传）

Windows、Python 3.12及系统黑体 C:/Windows/Fonts/simhei.ttf。
在解压后的 experiment3 目录安装 requirements.txt 依赖，再执行：

```
python build-report3.py --no-package
```

报告写回本目录，页面预览写入.preview。修改正文时先编辑content3.md。
本脚本不会访问GitHub、上传内容或读取项目密钥、数据库。
在项目原目录运行时沿用output/pdf布局；在本上传目录运行时自动使用当前目录。

## 实际状态

本地文件准备完成；远端是否提交以GitHub实际记录为准。实验2的提交编号不能作为实验3证据。
没有编码前原型记录，报告已说明三张原型属于当前设计整理；是否接受补做须由教师确认。
''', encoding="utf-8")
    files=sorted(p for p in target.rglob('*') if p.is_file() and p.name!='manifest.sha256')
    (target/'manifest.sha256').write_text(''.join(hashlib.sha256(p.read_bytes()).hexdigest()+'  '+p.relative_to(target).as_posix()+'\n' for p in files),encoding='utf-8')
    archive = OUT / "实验3-个人仓库上传包.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
        for source in package.rglob("*"):
            if source.is_file(): zf.write(source, source.relative_to(package))
    return package, archive


def main():
    parser=argparse.ArgumentParser(description='重建实验3报告与两张修订图件')
    parser.add_argument('--no-package',action='store_true',help='不生成上传包')
    args=parser.parse_args()
    draw_architecture()
    draw_flow()
    report = prepare_markdown()
    blocks = parse_blocks(report)
    build_pdf(blocks); build_docx(blocks)
    if QA.exists():
        for old in QA.glob("page-*.png"):
            old.unlink()
    QA.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(str(REPORT_PDF))
    for index, page in enumerate(pdf, 1): page.render(scale=1.2).to_pil().save(QA / f"page-{index:02}.png")
    page_count = len(pdf); pdf.close()
    result={"pdf": str(REPORT_PDF), "docx": str(REPORT_DOCX), "pages": page_count}
    if not PACKED and not args.no_package:
        package, archive = make_upload_package()
        result.update(upload_package=str(archive),package_dir=str(package))
    print(result)


if __name__ == "__main__": main()
